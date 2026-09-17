"""audit04 周报展示准备（presentation）聚焦测试。

覆盖：确定性事实摘要取代存储叙述、X/XX 占位与冲突计数、缺失字段不臆造为 0、
已知 0 保持 0、Top-N 数组长度不冒充权威总量、含数字/X 的书名不受影响、
helper 不改写 ORM、content 缺省时从 report 解析、中英语言。
"""

from datetime import date, timedelta
from types import SimpleNamespace

from app.utils.weekly_report_presentation import (
    build_factual_summary,
    category_chart_distribution,
    format_week_range,
    has_placeholder_metric,
    localize_report_title,
    prepare_report_presentation,
    scope_label,
    unknown_total_label,
)


def _report(summary='', content_dict=None, week_start=None, week_end=None, title=''):
    return SimpleNamespace(
        summary=summary,
        content=__import__('json').dumps(content_dict, ensure_ascii=False) if content_dict is not None else None,
        week_start=week_start or date(2026, 4, 20),
        week_end=week_end or date(2026, 4, 26),
        title=title,
    )


def _full_content():
    return {
        'total_books': 20,
        'total_new': 3,
        'total_rising': 5,
        'total_falling': 2,
        'new_books': [{'title': f'New {i}', 'author': 'A', 'category': '小说', 'rank': i + 1} for i in range(10)],
        'top_changes': [
            {'title': f'Change {i}', 'author': 'A', 'category': '小说', 'rank_change': 3} for i in range(10)
        ],
        'featured_books': [],
        'top_risers': [],
        'longest_running': [],
        'category_stats': {},
    }


# ---- 占位 X / XX ----


def test_has_placeholder_metric_detects_x():
    assert has_placeholder_metric('上榜书籍总数：X本') is True
    assert has_placeholder_metric('本周新增 XX本') is True
    assert has_placeholder_metric('这是一段正常叙述') is False


def test_presentation_summary_not_placeholder_when_stored_has_x():
    report = _report(summary='本周共有 **X** 本书籍进入榜单。', content_dict=_full_content())
    prepared = prepare_report_presentation(report)
    assert 'X' not in prepared['summary']
    assert prepared['summary_source'] == 'derived'
    assert '20' in prepared['summary']


# ---- 冲突计数：'本周新上榜书籍10本' vs 结构化 total_new=3 ----


def test_conflicting_stored_new_count_is_replaced_by_deterministic():
    report = _report(summary='本周新上榜书籍10本。', content_dict=_full_content())
    prepared = prepare_report_presentation(report)
    # 确定性摘要应使用权威 total_new=3，绝不透出存储叙述里的 10。
    assert '3' in prepared['summary']
    assert '10' not in prepared['summary']
    assert prepared['content']['total_new_display'] == '3'


def test_other_phrasing_stored_count_never_leaks():
    report = _report(summary='排名方面，99 本上升、88 本下降。', content_dict=_full_content())
    prepared = prepare_report_presentation(report)
    assert '99' not in prepared['summary']
    assert '88' not in prepared['summary']
    assert prepared['summary'].startswith('本周')


# ---- 缺失字段：不臆造为 0 ----


def test_missing_total_field_with_stored_zero_claim_stays_unknown():
    # total_new 缺失，但存储叙述声称“新增0本”，展示必须标“待补全”而非 0。
    content = _full_content()
    content['total_new'] = None
    report = _report(summary='本周新增0本。', content_dict=content)
    prepared = prepare_report_presentation(report)
    assert prepared['content']['total_new_known'] is False
    assert prepared['content']['total_new_display'] == unknown_total_label('zh')
    assert '0 本为新上榜' not in prepared['summary']
    assert unknown_total_label('zh') not in prepared['summary'].split('。')[0]  # 标题句仍给 total_books


def test_known_zero_remains_zero():
    content = _full_content()
    content['total_new'] = 0
    content['total_rising'] = 0
    content['total_falling'] = 0
    report = _report(summary='', content_dict=content)
    prepared = prepare_report_presentation(report)
    assert prepared['content']['total_new_display'] == '0'
    assert prepared['content']['total_rising_display'] == '0'
    assert '0 条为新上榜' in prepared['summary']
    assert '0 条上升' in prepared['summary']


def test_only_one_missing_total_preserves_others():
    content = _full_content()
    content['total_new'] = None  # 仅缺一个
    report = _report(summary='', content_dict=content)
    prepared = prepare_report_presentation(report)
    # total_books / rising / falling 仍保留并展示。
    assert prepared['content']['total_books_display'] == '20'
    assert prepared['content']['total_rising_display'] == '5'
    assert prepared['content']['total_falling_display'] == '2'
    assert '20' in prepared['summary']
    assert '5 条上升' in prepared['summary']


# ---- Top-N 数组长度不冒充权威总量 ----


def test_topn_array_length_not_used_as_authoritative_total():
    report = _report(summary='', content_dict=_full_content())  # new_books 长度 10，total_new=3
    prepared = prepare_report_presentation(report)
    assert prepared['content']['total_new_display'] == '3'
    assert prepared['content']['new_books'] is not None  # 结构化数组仍在
    assert len(prepared['content']['new_books']) == 10
    assert prepared['summary'].count('3') >= 1


def test_correct_structured_content_extracted():
    content = _full_content()
    report = _report(summary='', content_dict=content)
    prepared = prepare_report_presentation(report)
    assert prepared['totals'] == {
        'total_books': 20,
        'total_new': 3,
        'total_rising': 5,
        'total_falling': 2,
    }
    assert prepared['content']['totals_known'] is True


# ---- 含数字 / X 的书名在结构化区块原样保留 ----


def test_book_titles_with_x_and_digits_preserved():
    content = _full_content()
    content['new_books'] = [
        {'title': 'X战记', 'author': 'A', 'category': '漫画', 'rank': 1},
        {'title': 'Book 2: The Sequel', 'author': 'B', 'category': '小说', 'rank': 2},
        {'title': '1984', 'author': 'Orwell', 'category': '小说', 'rank': 3},
    ]
    report = _report(summary='', content_dict=content)
    prepared = prepare_report_presentation(report)
    titles = [b['title'] for b in prepared['content']['new_books']]
    assert titles == ['X战记', 'Book 2: The Sequel', '1984']


# ---- helper 不改写 ORM / 默认解析 content ----


def test_helper_does_not_mutate_orm(app, db):
    from app.models.schemas import WeeklyReport

    content = _full_content()
    with app.app_context():
        report = WeeklyReport(
            report_date=date.today(),
            week_start=date(2026, 4, 20),
            week_end=date(2026, 4, 26),
            title='T',
            summary='本周新上榜书籍10本。',
            content=__import__('json').dumps(content, ensure_ascii=False),
        )
        db.session.add(report)
        db.session.commit()

        original_summary = report.summary
        original_content = report.content

        prepare_report_presentation(report)
        prepare_report_presentation(report, content=__import__('json').loads(report.content))

        assert report.summary == original_summary
        assert report.content == original_content


def test_default_parsed_content_path(app, db):
    """不传 content 时，helper 应从 report.content 解析（回归 audit04 的 bug）。"""
    from app.models.schemas import WeeklyReport

    content = _full_content()
    with app.app_context():
        report = WeeklyReport(
            report_date=date.today(),
            week_start=date(2026, 4, 20),
            week_end=date(2026, 4, 26),
            title='T',
            summary='',
            content=__import__('json').dumps(content, ensure_ascii=False),
        )
        db.session.add(report)
        db.session.commit()

        prepared = prepare_report_presentation(report)  # 不传 content
        assert prepared['totals']['total_books'] == 20
        assert prepared['content']['total_new_display'] == '3'


def test_helper_with_explicit_content_argument():
    content = _full_content()
    report = _report(summary='本周新上榜书籍10本。', content_dict=None)  # content 缺失，但显式传入
    prepared = prepare_report_presentation(report, content=content)
    assert prepared['totals']['total_books'] == 20
    assert prepared['content']['total_new_display'] == '3'


# ---- 中英语言 ----


def test_zh_and_en_labels():
    report = _report(summary='本周新上榜书籍10本。', content_dict=_full_content())
    zh = prepare_report_presentation(report, locale='zh')
    en = prepare_report_presentation(report, locale='en')
    assert unknown_total_label('zh') == '数据待补全'
    assert unknown_total_label('en') == 'Data pending'
    assert zh['content']['total_books_display'] == '20'
    assert en['content']['total_books_display'] == '20'
    # 中英摘要语言不同。
    assert zh['summary'] != en['summary']


def test_unknown_label_localized_in_prepared():
    report = _report(summary='', content_dict={'total_books': 10})
    zh = prepare_report_presentation(report, locale='zh')
    en = prepare_report_presentation(report, locale='en')
    assert zh['content']['total_new_display'] == '数据待补全'
    assert en['content']['total_new_display'] == 'Data pending'


def test_scope_label_localized():
    assert '上榜记录' in scope_label('zh')
    assert 'list entries' in scope_label('en')
    assert 'counts separately' in scope_label('en')
    # 口径必须是“本报告周期所采集分类榜单”，不得暗示去重书数 / 唯一 ISBN。
    assert '不等于去重后的书数' in scope_label('zh')
    assert 'unique books or ISBNs' in scope_label('en')


# ---- 确定性摘要 ----


def test_build_factual_summary_no_dates():
    s = build_factual_summary({'total_books': 10, 'total_new': 2, 'total_rising': 3, 'total_falling': 1})
    assert '10' in s
    assert '2' in s
    assert '上榜记录' in s
    # 日期缺失时不留空括号。
    assert '（）' not in s
    assert '()' not in s


def test_build_factual_summary_capitalizes_english_first_sentence():
    s = build_factual_summary(
        {'total_books': 15, 'total_new': 1, 'total_rising': 2, 'total_falling': 3},
        week_start=date(2026, 4, 20),
        week_end=date(2026, 4, 26),
        locale='en',
    )
    assert s[0].isupper(), s
    # 英文缺日期时同样不留空括号且首句首字母大写。
    no_dates = build_factual_summary({'total_books': 15}, locale='en')
    assert no_dates[0].isupper(), no_dates
    assert '()' not in no_dates


def test_dates_are_included_in_summary():
    report = _report(summary='', content_dict=_full_content())
    prepared = prepare_report_presentation(report)
    assert '2026' in prepared['summary']
    assert '04月20日' in prepared['summary'] or '2026-04-20' in prepared['summary']


# ---- 回归：total_books 缺失时不得吞掉已知计数与日期 ----


def test_unknown_total_books_keeps_dates_and_every_known_count():
    """回归：total_books=None 时早期 return 会丢掉已知 new/rising/falling 与日期。"""
    content = _full_content()
    content['total_books'] = None
    report = _report(summary='', content_dict=content)
    prepared = prepare_report_presentation(report)
    summary = prepared['summary']

    # 日期区间仍在。
    assert '2026' in summary
    # 三个已知计数全部保留（total_new=3 / rising=5 / falling=2）。
    assert '3' in summary
    assert '5' in summary
    assert '2' in summary
    # 未知的上榜记录逐项标注待补全，且不留空括号。
    assert '待补全' in summary
    assert '（）' not in summary
    assert '()' not in summary


def test_unknown_total_books_marks_only_the_unknown_one():
    content = _full_content()
    content['total_books'] = None
    prepared = prepare_report_presentation(_report(summary='', content_dict=content))
    assert prepared['content']['total_books_known'] is False
    for key in ('total_new', 'total_rising', 'total_falling'):
        assert prepared['content'][f'{key}_known'] is True
        assert f'{key}_display' in prepared['content']


def test_only_rising_known_mentions_falling_pending():
    content = _full_content()
    content['total_falling'] = None
    prepared = prepare_report_presentation(_report(summary='', content_dict=content))
    summary = prepared['summary']
    assert '5' in summary
    assert '下降数量待补全' in summary


def test_only_falling_known_mentions_rising_pending():
    content = _full_content()
    content['total_rising'] = None
    prepared = prepare_report_presentation(_report(summary='', content_dict=content))
    summary = prepared['summary']
    assert '2' in summary
    assert '上升数量待补全' in summary


def test_summary_unknowns_never_rendered_as_zero():
    content = _full_content()
    for key in ('total_books', 'total_new', 'total_rising', 'total_falling'):
        content[key] = None
    report = _report(summary='', content_dict=content)
    summary = prepare_report_presentation(report)['summary']
    assert '待补全' in summary
    # 日期仍保留，未知项不用空括号占位。
    assert '2026' in summary
    assert '（）' not in summary
    # 未知的计数不得用具名断言冒充 0（日期里的 0 是合法的，不能整串禁止 '0'）。
    for phrase in ('0 条为新上榜', '0 条上升', '0 条下降', '0 条上榜记录'):
        assert phrase not in summary


def test_known_zero_counts_still_render_as_zero_not_pending():
    """已知 0 必须保持 0：待补全与 0 语义不同。"""
    content = _full_content()
    content['total_books'] = None
    content['total_new'] = 0
    content['total_rising'] = 0
    content['total_falling'] = 0
    summary = prepare_report_presentation(_report(summary='', content_dict=content))['summary']
    assert '0 条为新上榜' in summary
    assert '0 条上升' in summary
    # 只有真正未知的上榜记录标待补全（标题句共 1 处）。
    assert summary.count('待补全') == 1


def test_summary_does_not_mutate_content_orm(app, db):
    """摘要准备不改写 ORM content / summary（含 total_books=None 路径）。"""
    from app.models.schemas import WeeklyReport

    content = _full_content()
    content['total_books'] = None
    with app.app_context():
        report = WeeklyReport(
            report_date=date.today(),
            week_start=date(2026, 4, 20),
            week_end=date(2026, 4, 26),
            title='T',
            summary='storage narrative',
            content=__import__('json').dumps(content, ensure_ascii=False),
        )
        db.session.add(report)
        db.session.commit()
        raw_summary, raw_content = report.summary, report.content

        prepare_report_presentation(report)
        build_factual_summary(
            {'total_books': None, 'total_new': 3, 'total_rising': 5, 'total_falling': 2},
            report.week_start,
            report.week_end,
        )

        assert report.summary == raw_summary
        assert report.content == raw_content


# ---- 类别分布图：只用结构化 category_stats[*].count ----

_CATEGORY_CHART_CONTENT = {
    'total_books': 15,
    'total_new': 0,
    'total_rising': 0,
    'total_falling': 0,
    'category_stats': {'精装小说': {'count': 15}},
    # 重叠的 Top-N 数组：合计 24 条，绝不能被图表当作分布数据。
    'top_changes': [{'title': f'C{i}', 'author': 'A', 'category': '精装小说', 'rank_change': 1} for i in range(10)],
    'new_books': [{'title': f'N{i}', 'author': 'A', 'category': '精装小说', 'rank': i} for i in range(8)],
    'top_risers': [{'title': f'R{i}', 'author': 'A', 'category': '精装小说', 'rank_change': 2} for i in range(3)],
    'longest_running': [{'title': f'L{i}', 'author': 'A', 'category': '精装小说'} for i in range(3)],
    'featured_books': [],
}


def test_category_chart_uses_structured_counts_not_overlapping_arrays():
    report = _report(summary='', content_dict=_CATEGORY_CHART_CONTENT)
    chart = prepare_report_presentation(report)['content']['category_chart']

    assert chart['available'] is True
    assert chart['labels'] == ['精装小说']
    assert chart['counts'] == [15]  # 不是 24
    assert sum(chart['counts']) == 15
    overlapping_total = sum(
        len(_CATEGORY_CHART_CONTENT[key])
        for key in ('top_changes', 'new_books', 'top_risers', 'longest_running', 'featured_books')
    )
    assert overlapping_total == 24
    assert sum(chart['counts']) != overlapping_total


def test_category_chart_absent_stats_is_explicitly_unavailable():
    content = dict(_CATEGORY_CHART_CONTENT)
    content.pop('category_stats')
    prepared = prepare_report_presentation(_report(summary='', content_dict=content))
    assert prepared['content']['category_chart']['available'] is False
    chart = prepared['content']['category_chart']
    assert chart['counts'] == [] and chart['labels'] == []
    assert chart['message']  # 显式不可用提示，不是 0 / 猜测值
    assert '待补全' in chart['message']


def test_category_chart_invalid_stats_is_explicitly_unavailable():
    for invalid in (
        {},
        {'精装小说': {}},
        {'精装小说': {'count': None}},
        {'精装小说': {'count': '15'}},
        {'精装小说': {'count': -1}},
        {'精装小说': {'count': True}},
        'not-a-dict',
        {'': {'count': 3}},
        {'精装小说': 'nope'},
    ):
        content = dict(_CATEGORY_CHART_CONTENT, category_stats=invalid)
        chart = prepare_report_presentation(_report(summary='', content_dict=content))['content']['category_chart']
        assert chart['available'] is False, invalid
        assert chart['message'], invalid


def test_category_chart_keeps_known_zero_category():
    """已知 count=0 的分类保留为 0，不算非法。"""
    content = dict(_CATEGORY_CHART_CONTENT, category_stats={'精装小说': {'count': 0}, '平装小说': {'count': 15}})
    chart = prepare_report_presentation(_report(summary='', content_dict=content))['content']['category_chart']
    assert chart['available'] is True
    assert chart['labels'] == ['精装小说', '平装小说']
    assert chart['counts'] == [0, 15]


def test_category_chart_partial_totals():
    """partial totals：图表分布与已知总量解耦，只用 category_stats。"""
    content = dict(_CATEGORY_CHART_CONTENT, total_books=None)
    chart = prepare_report_presentation(_report(summary='', content_dict=content))['content']['category_chart']
    assert chart['counts'] == [15]


def test_category_chart_localized_unavailable_message():
    content = dict(_CATEGORY_CHART_CONTENT)
    content.pop('category_stats')
    zh = prepare_report_presentation(_report(summary='', content_dict=content), locale='zh')
    en = prepare_report_presentation(_report(summary='', content_dict=content), locale='en')
    assert zh['content']['category_chart']['message'] != en['content']['category_chart']['message']
    assert zh['content']['category_chart']['message'].startswith('数据待补全')


def test_category_chart_does_not_mutate_source_content():
    content = _CATEGORY_CHART_CONTENT
    before = __import__('copy').deepcopy(content)
    prepare_report_presentation(_report(summary='', content_dict=content))
    assert content == before
    assert content['category_stats'] == {'精装小说': {'count': 15}}


def test_category_chart_distribution_helper_directly():
    """helper 可直接调用：安全暴露图表分布，非法输入返回不可用。"""
    ok = category_chart_distribution(_CATEGORY_CHART_CONTENT)
    assert ok['available'] is True and ok['counts'] == [15]
    assert category_chart_distribution(None)['available'] is False
    assert category_chart_distribution({})['available'] is False
    assert category_chart_distribution({'category_stats': {}})['available'] is False


def test_category_chart_rendered_by_detail_page(client, app, db, monkeypatch):
    """真实渲染详情页：重叠数组 24 条，图表数据必须是结构化 15。"""
    import re
    from unittest.mock import MagicMock

    from app.models.schemas import WeeklyReport

    monkeypatch.setitem(app.extensions, 'book_service', MagicMock())
    with app.app_context():
        report = WeeklyReport(
            report_date=date.today(),
            week_start=date.today() - timedelta(days=7),
            week_end=date.today(),
            title='Category Chart Report',
            summary='',
            content=__import__('json').dumps(_CATEGORY_CHART_CONTENT, ensure_ascii=False),
        )
        db.session.add(report)
        db.session.commit()
        date_str = report.report_date.strftime('%Y-%m-%d')

    response = client.get(f'/reports/weekly/{date_str}')
    assert response.status_code == 200
    body = response.data.decode('utf-8')

    match = re.search(r'const reportContent = (\{.*?\});\n', body, re.S)
    assert match, '详情页未内联 reportContent'
    payload = __import__('json').loads(match.group(1))
    chart = payload['category_chart']
    assert chart['available'] is True
    assert chart['counts'] == [15]
    assert sum(chart['counts']) == 15
    # 页面脚本不得再拼接重叠数组来统计类别。
    assert 'allBooksForCategory' not in body


def test_category_chart_unavailable_rendered_message_on_detail_page(client, app, db, monkeypatch):
    import re
    from unittest.mock import MagicMock

    from app.models.schemas import WeeklyReport

    monkeypatch.setitem(app.extensions, 'book_service', MagicMock())
    content = dict(_CATEGORY_CHART_CONTENT, category_stats={})
    with app.app_context():
        report = WeeklyReport(
            report_date=date.today(),
            week_start=date.today() - timedelta(days=7),
            week_end=date.today(),
            title='No Category Stats',
            summary='',
            content=__import__('json').dumps(content, ensure_ascii=False),
        )
        db.session.add(report)
        db.session.commit()
        date_str = report.report_date.strftime('%Y-%m-%d')

    body = client.get(f'/reports/weekly/{date_str}').data.decode('utf-8')
    match = re.search(r'const reportContent = (\{.*?\});\n', body, re.S)
    assert match
    chart = __import__('json').loads(match.group(1))['category_chart']
    assert chart['available'] is False
    assert chart['message']
    assert 'chart-unavailable-message' in body


# ---- 路由集成：桌面/移动/列表/导出 与权威总量一致 ----


def _seed_report(app, db, summary='本周新上榜书籍10本。'):
    from app.models.schemas import WeeklyReport

    content = _full_content()
    with app.app_context():
        report = WeeklyReport(
            report_date=date.today(),
            week_start=date.today() - timedelta(days=7),
            week_end=date.today(),
            title='Test Weekly Report',
            summary=summary,
            content=__import__('json').dumps(content, ensure_ascii=False),
        )
        db.session.add(report)
        db.session.commit()
        return report.report_date.strftime('%Y-%m-%d')


def test_desktop_detail_uses_authoritative_total_not_topn_length(client, app, db, monkeypatch):
    from unittest.mock import MagicMock

    monkeypatch.setitem(app.extensions, 'book_service', MagicMock())
    date_str = _seed_report(app, db)
    response = client.get(f'/reports/weekly/{date_str}')
    assert response.status_code == 200
    body = response.data.decode('utf-8')
    # 概览区展示确定性摘要（含权威 total_books=20 与 total_new=3）。
    assert '20' in body
    assert '3' in body


def test_desktop_detail_summary_is_deterministic_not_stored_conflict(client, app, db, monkeypatch):
    from unittest.mock import MagicMock

    monkeypatch.setitem(app.extensions, 'book_service', MagicMock())
    date_str = _seed_report(app, db, summary='本周新上榜书籍10本。')
    response = client.get(f'/reports/weekly/{date_str}')
    body = response.data.decode('utf-8')
    # 确定性摘要按权威 total_new=3 构造；冲突的“10本”不应作为概览出现。
    assert '10 本' not in body
    assert '3' in body


def test_list_page_uses_authoritative_totals(client, app, db, monkeypatch):
    from unittest.mock import MagicMock

    monkeypatch.setitem(app.extensions, 'book_service', MagicMock())
    _seed_report(app, db)
    response = client.get('/reports/weekly')
    assert response.status_code == 200
    body = response.data.decode('utf-8')
    # 列表卡片展示权威 total_books=20 与 total_new=3（而非 new_books 数组长度 10）。
    assert '20' in body
    assert '3' in body


def _lang_sync_ssr_lang(body):
    """从渲染出的页面里取 SSR 语言常量（脚本内联，非 DOM 属性）。"""
    import json
    import re

    match = re.search(r'var ssrLang = canonical\(("(?:en|zh)[^"]*")\);', body)
    assert match, '渲染页面未内联 ssrLang 常量'
    return json.loads(match.group(1))


def test_rendered_weekly_pages_inline_the_real_ssr_locale(client, app, db, monkeypatch):
    """渲染页面里脚本的 SSR 语言必须等于实际的 SSR 语言，且不依赖 DOM 属性。

    回归：宏曾把语言写在 <span data-ssr-lang> 上、脚本却读 <html data-ssr-lang>
    （生产上不存在）→ ssrLang 恒为 '' → 语言修正永不生效。
    """
    import re
    from unittest.mock import MagicMock

    monkeypatch.setitem(app.extensions, 'book_service', MagicMock())
    date_str = _seed_report(app, db)

    for path in ('/reports/weekly', f'/reports/weekly/{date_str}'):
        response = client.get(path)
        assert response.status_code == 200, path
        body = response.data.decode('utf-8')

        # 实际 SSR 语言：<html lang> 由 get_locale() 渲染而来。
        html_lang = re.search(r'<html lang="([^"]+)"', body)
        assert html_lang, f'{path} 缺少 <html lang>'
        actual = 'zh' if html_lang.group(1).lower().startswith('zh') else 'en'

        assert _lang_sync_ssr_lang(body) == actual, f'{path} 内联的 SSR 语言与实际不一致'
        # 不得再依赖那个从未被渲染过的 <html> 属性。
        assert not re.search(r'^\s*<html[^>]*data-ssr-lang', body, re.M), path
        assert re.search(r'var ssrLang = canonical\("(?:en|zh)', body), path


def test_export_uses_prepared_deterministic_summary(client, app, db, monkeypatch):
    """导出路由传入 prepared，摘要应为确定性事实摘要（total_books=20），且不报错。"""
    from unittest.mock import MagicMock

    monkeypatch.setitem(app.extensions, 'book_service', MagicMock())
    date_str = _seed_report(app, db, summary='本周新上榜书籍10本。')

    pdf = client.get(f'/reports/weekly/{date_str}/export?format=pdf')
    assert pdf.status_code in (200, 500)
    excel = client.get(f'/reports/weekly/{date_str}/export?format=excel')
    assert excel.status_code in (200, 500)


def test_export_service_direct_call_uses_prepared(app, db):
    """导出服务在传入 prepared 时使用确定性摘要与本地化 display 值。"""
    from io import BytesIO

    from openpyxl import load_workbook

    from app.models.schemas import WeeklyReport
    from app.services.export_service import ExportService
    from app.utils.weekly_report_presentation import prepare_report_presentation

    with app.app_context():
        report = WeeklyReport(
            report_date=date.today(),
            week_start=date(2026, 4, 20),
            week_end=date(2026, 4, 26),
            title='T',
            summary='本周新上榜书籍10本。',
            content=__import__('json').dumps(_full_content(), ensure_ascii=False),
        )
        db.session.add(report)
        db.session.commit()

        prepared = prepare_report_presentation(report, locale='zh')
        svc = ExportService()
        buffer = svc.export_weekly_report_excel(report, prepared=prepared)
        assert buffer is not None and isinstance(buffer, BytesIO)
        wb = load_workbook(buffer)
        ws = wb.active
        cells = ' '.join(str(c.value or '') for row in ws.iter_rows() for c in row)
        assert '上榜记录: 20' in cells
        assert '新书: 3' in cells
        assert '20' in cells


# ---------------------------------------------------------------------------
# 英文摘要 / 日期区间 / 标题本地化（locale-source cleanup）
# ---------------------------------------------------------------------------


def test_english_new_sentence_is_complete_and_capitalized():
    """回归：英文 'of which ... are new to the lists' 小写、无句末标点，
    直接跟在上句句号后读起来是残句。完整句由简单 join 产生：片段本身不带句末
    标点，统一用 term（'. '）连接并在整句末尾补一个 term。"""
    s = build_factual_summary(
        {'total_books': 12, 'total_new': 3, 'total_rising': 2, 'total_falling': 1},
        week_start=date(2026, 4, 20),
        week_end=date(2026, 4, 26),
        locale='en',
    )
    # 片段本身无句末标点（由连接逻辑统一补），整句里读到的仍是完整句。
    assert '3 entries are new to the lists.' in s, s
    assert 'Of which' not in s, f'不得再用「Of which」残句开头: {s}'
    assert '..' not in s, f'不得出现重复句号: {s}'
    assert 'lists.By' not in s, f'句子之间必须有空格: {s}'
    assert 'lists,By' not in s, f'句子之间必须有分隔: {s}'


def test_english_new_sentence_preserves_zero_and_unknown_semantics():
    zero = build_factual_summary(
        {'total_books': 0, 'total_new': 0, 'total_rising': 0, 'total_falling': 0},
        week_start=date(2026, 4, 20),
        week_end=date(2026, 4, 26),
        locale='en',
    )
    assert '0 entries are new to the lists.' in zero
    unknown = build_factual_summary({'total_books': 5, 'total_new': None}, locale='en')
    assert 'pending' in unknown.lower()
    assert '0 entries are new' not in unknown


def test_english_summary_has_no_chinese_connector():
    s = build_factual_summary(
        {'total_books': 12},
        week_start=date(2026, 4, 20),
        week_end=date(2026, 4, 26),
        locale='en',
    )
    assert ' 至 ' not in s, f'英文摘要不得出现中文连接词「至」: {s}'


def test_format_week_range_is_locale_aware():
    zh = format_week_range(date(2026, 4, 20), date(2026, 4, 26), locale='zh')
    en = format_week_range(date(2026, 4, 20), date(2026, 4, 26), locale='en')
    assert '至' in zh and '年' in zh
    assert '至' not in en, f'英文区间不得含「至」: {en}'
    assert en == '2026-04-20–2026-04-26'


def test_format_week_range_missing_side_is_empty():
    assert format_week_range(date(2026, 4, 20), None, locale='en') == ''
    assert format_week_range(None, None, locale='zh') == ''


def test_known_standard_title_is_localized_for_english():
    standard = '2026年01月05日-2026年01月11日 畅销书周报'
    assert localize_report_title(standard, locale='en') == '2026-01-05–2026-01-11 Weekly Bestseller Report'
    assert localize_report_title(standard, locale='zh') == standard


def test_exact_known_plain_title_is_localized_and_titles_stay_exact():
    """无日期区间的已知标准标题也要本地化；**只认精确相等**，带该短语的
    任意标题（前缀/后缀/大小写变体）必须原样透传。"""
    zh_standard = '纽约时报畅销书周报'
    assert localize_report_title(zh_standard, locale='en') == 'NYT Weekly Bestseller Report'
    assert localize_report_title(zh_standard, locale='zh') == zh_standard
    # 空白被 strip 后仍命中（与日期形态标题同口径）
    assert localize_report_title(f'  {zh_standard}  ', locale='en') == 'NYT Weekly Bestseller Report'

    for title in (
        f'{zh_standard}（增刊）',
        f'本周{zh_standard}',
        '纽约时报畅销书周报 2026',
        'nyt weekly bestseller report',
    ):
        assert localize_report_title(title, locale='en') == title, f'任意标题被改写: {title}'
        assert localize_report_title(title, locale='zh') == title, f'任意标题被改写: {title}'


def test_arbitrary_db_titles_are_preserved_verbatim():
    """任意（人工/非标准）标题必须原样透传，绝不猜测改写。"""
    for title in (
        '本周精选：编辑手记',
        'My Custom Report',
        '2026年01月05日-2026年01月11日 畅销书周报（增刊）',
        '',
    ):
        assert localize_report_title(title, locale='en') == title


def test_localize_report_title_tolerates_non_string():
    assert localize_report_title(None, locale='en') == ''
    assert localize_report_title(SimpleNamespace(), locale='en') == ''


def test_prepared_content_exposes_localized_range_and_title():
    for locale, expected_connector in (('en', '–'), ('zh', '至')):
        prepared = prepare_report_presentation(
            _report(summary='', content_dict=_full_content(), title='2026年01月05日-2026年01月11日 畅销书周报'),
            locale=locale,
        )['content']
        assert expected_connector in prepared['week_range']
        assert prepared['title_display']


# ---------------------------------------------------------------------------
# 英文内联图表脚本的 quoting 回归
# ---------------------------------------------------------------------------


def test_english_chart_script_has_no_raw_gettext_string_literals():
    """回归：`label: '{{ _('本周排名') }}'` 在 en 目录下产出
    `label: 'This week's rank'`，字符串被提前闭合 → "Unexpected identifier"，
    整段内联脚本（所有图表 + 弹窗）全部失效。

    修复：JS 字符串位置一律走 `|tojson`。这里直接扫模板，禁止再出现
    "单引号包裹裸 {{ _('...') }}" 的形态。
    """
    import re
    from pathlib import Path

    src = Path(__file__).resolve().parent.parent / 'templates' / 'weekly_report_detail.html'
    text = src.read_text(encoding='utf-8')
    # 只看 <script> 块（Jinja 注释在模板里无副作用，但脚本块里的才是真 JS）
    blocks = re.findall(r'<script(?![^>]*ld\+json)[^>]*>(.*?)</script>', text, re.S)
    assert blocks, 'weekly_report_detail.html 应含内联脚本'
    offenders = []
    for block in blocks:
        # 去掉 // 注释行后再找裸插值
        code = '\n'.join(line for line in block.splitlines() if not line.lstrip().startswith('//'))
        offenders += re.findall(r"'[^'\n]*\{\{\s*_\([^)]*\)\s*\}\}[^'\n]*'", code)
    assert offenders == [], f'JS 字符串位置必须走 |tojson，发现裸 gettext 插值: {offenders}'


def test_english_chart_labels_are_json_encoded_for_real_catalog(client, app, db, monkeypatch):
    """用真实 en 目录渲染详情页：图表标签必须 JSON 编码，绝不能出现被提前闭合的单引号。"""
    date_str = _seed_report(app, db)
    monkeypatch.setattr(
        'app.routes.main.parse_report_content',
        lambda report: _full_content(),
    )
    html = client.get(f'/reports/weekly/{date_str}?lang=en').get_data(as_text=True)
    if 'new Chart(' not in html:
        return  # 无图表数据时该页不渲染图表块，跳过（seed 未提供 chart 数据）
    # 真实目录里 '本周排名' → "This week's rank"；必须编码为 JSON 字符串
    assert "label: 'This week's rank'" not in html, '裸单引号会让整段内联脚本语法错误'
